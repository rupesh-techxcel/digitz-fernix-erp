# Copyright (c) 2026, Rupesh P and contributors
# For license information, please see license.txt

"""Day Open / Day Close for a cashier at a counter.

A Counter Session is one cashier's day at one registered counter PC: opened with a
float (open_day), closed by counting the till (close_day). Every cash document a
cashier saves -- Sales Invoice, Sales Return, Receipt Entry, Expense Entry -- is
stamped with the open session (stamp_counter_session), which is how the session
works out the cash it should hold. A cashier without an open day on this counter
cannot save them; supervisors and background jobs (the token sync) are not held to it.

Shift change on one counter: the outgoing cashier closes with Close & Hand Over
(close_day, then log out), and the next cashier opens with the counted cash as
their float (get_last_close carries the count). A day left open by a cashier who
has gone is closed by a supervisor who counts the till (supervisor_close).
"""

import json

import frappe
from frappe.utils import add_days, cint, date_diff, flt, getdate, now_datetime, nowdate

from digitz_erp.api.counter_api import get_request_device, get_request_registration, is_supervisor
from digitz_erp.digitz_erp.doctype.counter_session.counter_session import TOLERANCE

# The day's sales and what the till should hold: shown to supervisors only. The
# same fields are permlevel 1 on Counter Session.
SUPERVISOR_ONLY_FIELDS = ("cash_sales", "cash_receipts", "cash_refunds", "cash_paid_out", "expected_cash", "difference")

# The note and coin values offered in the Close Day cash count (AED)
DENOMINATIONS = (1000, 500, 200, 100, 50, 20, 10, 5, 1, 0.5, 0.25)


def get_open_session(user=None, counter=None, billing=False):
	"""The user's unsubmitted session (optionally on `counter`), as a name, or None.

	With `billing`, only a session still Open counts: one waiting in Closing for a
	supervisor takes no more documents.
	"""
	filters = {"cashier": user or frappe.session.user, "docstatus": 0}
	if counter:
		filters["counter"] = counter
	if billing:
		filters["status"] = "Open"
	return frappe.db.get_value("Counter Session", filters, "name")


def must_have_open_day(user=None):
	user = user or frappe.session.user
	return "Cashier" in frappe.get_roles(user) and not is_supervisor(user)


def stamp_counter_session(doc):
	"""Record the counter PC and the open day a cash document is saved in.

	Called from before_validate of Sales Invoice, Sales Return, Receipt Entry and
	Expense Entry. Sets `counter_session` (and `counter` / `counter_device` where
	the doctype has them) from the request's registered PC and the user's open
	session on it. A cashier who is not a supervisor is stopped without both.
	Nothing is required of background work, which has no request.
	"""
	if not getattr(frappe.local, "request", None):
		return

	device = get_request_device()
	if device:
		if doc.meta.has_field("counter"):
			doc.counter = device.counter
		if doc.meta.has_field("counter_device"):
			doc.counter_device = device.device

	session = get_open_session(counter=device.counter, billing=True) if device else None
	if session:
		doc.counter_session = session
		return

	if not must_have_open_day():
		return

	if not device:
		frappe.throw(
			"This device is not registered as a counter, so this cannot be saved here. "
			"Ask a supervisor to register it.",
			title="Unregistered Device",
		)

	frappe.throw(
		f"Open the day on {device.counter} before billing. Click the counter badge at the top of the page.",
		title="Day Not Opened",
	)


# How far back a cashier may date a Sales Invoice, or correct a submitted one
CASHIER_BACKDATE_DAYS = 10


def cashier_date_window():
	"""(earliest, latest) posting date a cashier may use: today and the days before it."""
	today = getdate(nowdate())
	return add_days(today, -CASHIER_BACKDATE_DAYS), today


def restrict_cashier_posting_date(doc):
	"""A cashier (not a supervisor) dates an invoice today or up to
	CASHIER_BACKDATE_DAYS back, never in the future -- except an invoice being
	corrected, which keeps the date it was submitted with.

	Called from Sales Invoice validate, so on every save and submit. Both the date
	saved before (an old draft) and the one being saved must be in the window:
	anything older is for a supervisor. Nothing is required of background work
	(the token sync), which has no request.
	"""
	if not getattr(frappe.local, "request", None) or not must_have_open_day():
		return

	# A submitted invoice the cashier took back to correct keeps its own date,
	# however old (sales_invoice_api.correct_sales_invoice)
	from digitz_erp.api.sales_invoice_api import correction_date

	kept = correction_date(doc)
	if kept and getdate(kept) == getdate(doc.posting_date):
		return

	earliest, today = cashier_date_window()
	# A correction moved to another date is held to the window by its new date only
	saved_date = None if doc.is_new() or kept else frappe.db.get_value(doc.doctype, doc.name, "posting_date")
	label = "This invoice" if doc.is_new() else doc.name

	for posting_date in (saved_date, doc.posting_date):
		if not posting_date:
			continue
		posting_date = getdate(posting_date)
		if posting_date > today:
			frappe.throw(
				f"{label} is dated {frappe.utils.formatdate(posting_date)}, in the future. "
				"A cashier cannot date an invoice after today.",
				title="Future Date",
			)
		if posting_date < earliest:
			frappe.throw(
				f"{label} is dated {frappe.utils.formatdate(posting_date)}. A cashier can date an invoice "
				f"at most {CASHIER_BACKDATE_DAYS} days back (from {frappe.utils.formatdate(earliest)}). "
				"Ask a supervisor to handle this one.",
				title="Date Too Far Back",
			)


def get_last_close(counter):
	"""The last counted session on `counter`, or None.

	{name, cashier, cashier_name, closed_on, counted_cash, difference, status,
	denominations}. A close still waiting for approval counts: its till was
	counted and handed over, so it is what the next day starts from.
	"""
	if not counter:
		return None
	rows = frappe.db.sql(
		"""
		SELECT name, cashier, closed_by, closed_on, counted_cash, difference, status
		FROM `tabCounter Session`
		WHERE counter = %(counter)s AND closed_on IS NOT NULL
		  AND (docstatus = 1 OR (docstatus = 0 AND status = 'Closing'))
		ORDER BY closed_on DESC
		LIMIT 1
		""",
		{"counter": counter},
		as_dict=True,
	)
	if not rows:
		return None
	last = rows[0]
	last.cashier_name = frappe.utils.get_fullname(last.cashier)
	last.closed_by_name = frappe.utils.get_fullname(last.closed_by) if last.closed_by else None
	last.denominations = frappe.get_all(
		"Cash Denomination Count",
		filters={"parent": last.name, "parenttype": "Counter Session", "parentfield": "denominations"},
		fields=["denomination", "count"],
		order_by="idx",
	)
	return last


@frappe.whitelist()
def get_state():
	"""Everything the counter badge shows: the PC, the user's day, and pending approvals."""
	device = get_request_device()
	supervisor = is_supervisor()

	session = None
	name = get_open_session()
	if name:
		session = frappe.db.get_value(
			"Counter Session", name, ["name", "counter", "status", "opened_on", "opening_float"], as_dict=True
		)

	return {
		"device": device.device if device else None,
		"counter": device.counter if device else None,
		"session": session,
		"can_register": supervisor,
		"is_supervisor": supervisor,
		"must_open_day": must_have_open_day(),
		"pending_approvals": frappe.db.count("Counter Session", {"docstatus": 0, "status": "Closing"}) if supervisor else 0,
		# PCs cashiers asked to register that would take over a counter
		"device_requests": frappe.db.count("Counter Device", {"approval_status": "Pending"}) if supervisor else 0,
		# This browser's own request, while it waits or after it was rejected
		"registration_request": get_request_registration() if not device else None,
		"is_cashier": "Cashier" in frappe.get_roles(),
		"counters": frappe.get_all("Counter", filters={"enabled": 1}, pluck="name", order_by="name") if supervisor else [],
		"denominations": DENOMINATIONS,
		# How far back a cashier may date or correct an invoice (restrict_cashier_posting_date)
		"backdate_days": CASHIER_BACKDATE_DAYS,
		# The non-cash modes Day Close asks a total for (card slips, bank transfers...)
		"payment_modes": get_non_cash_payment_modes(),
		"user_full_name": frappe.utils.get_fullname(frappe.session.user),
		"now": now_datetime(),
		# Shown when opening a day: the cash the last day on this counter closed with,
		# unless Settings asks for a blind count (then it is not sent at all)
		"show_last_close": show_last_close(),
		"last_close": get_last_close(device.counter) if device and not session and show_last_close() else None,
		# Another cashier's day still open on this counter, which stops a new one
		"counter_taken": get_counter_taken(device.counter) if device and not session else None,
	}


def show_last_close():
	"""Settings > Show Previous Day Close Balance in Day Open. Off (a blind count)
	unless a System Manager ticks it."""
	return cint(frappe.db.get_single_value("Settings", "show_previous_close_in_day_open"))


def get_counter_taken(counter):
	"""Someone else's Open day on `counter`: {name, cashier, cashier_name, opened_on}, or None."""
	row = frappe.db.get_value(
		"Counter Session",
		{"counter": counter, "docstatus": 0, "status": "Open", "cashier": ["!=", frappe.session.user]},
		["name", "cashier", "opened_on"],
		as_dict=True,
	)
	if row:
		row.cashier_name = frappe.utils.get_fullname(row.cashier)
	return row


def get_non_cash_payment_modes():
	return frappe.get_all(
		"Payment Mode", filters={"mode": ["!=", "Cash"]}, fields=["name as payment_mode", "mode"], order_by="name"
	)


def parse_payment_totals(payment_totals):
	"""[{payment_mode, amount}] for the non-cash modes given an amount, from the pages' JSON."""
	if isinstance(payment_totals, str):
		payment_totals = json.loads(payment_totals or "[]")
	non_cash = {m.payment_mode for m in get_non_cash_payment_modes()}
	rows = []
	for row in payment_totals or []:
		if row.get("amount") in (None, ""):
			continue
		if row.get("payment_mode") not in non_cash:
			frappe.throw(f"{row.get('payment_mode')} is not a non-cash payment mode.")
		if flt(row.get("amount")) < 0:
			frappe.throw(f"The {row.get('payment_mode')} total cannot be negative.")
		rows.append({"payment_mode": row.get("payment_mode"), "amount": flt(row.get("amount"))})
	return rows


def parse_denominations(denominations):
	"""[{denomination, count}] with counts, from the pages' JSON."""
	if isinstance(denominations, str):
		denominations = json.loads(denominations or "[]")
	return [
		{"denomination": flt(row.get("denomination")), "count": flt(row.get("count"))}
		for row in denominations or []
		if flt(row.get("count"))
	]


@frappe.whitelist()
def open_day(opening_float=0, denominations=None, remarks=None):
	"""Open the user's day on this PC's counter with the cash now in the till.

	`denominations` is the optional note-by-note count; when given, it makes up
	the opening float.
	"""
	device = get_request_device()
	if not device:
		frappe.throw("This device is not registered as a counter. Ask a supervisor to register it.")

	counted = parse_denominations(denominations)
	if counted:
		opening_float = sum(row["denomination"] * row["count"] for row in counted)

	if flt(opening_float) < 0:
		frappe.throw("The opening float cannot be negative.")

	last_close = get_last_close(device.counter)

	session = frappe.get_doc({
		"doctype": "Counter Session",
		"counter": device.counter,
		"counter_device": device.device,
		"cashier": frappe.session.user,
		"status": "Open",
		"opened_on": now_datetime(),
		"opening_float": flt(opening_float),
		"previous_session": last_close.name if last_close else None,
		"previous_closing_cash": flt(last_close.counted_cash) if last_close else 0,
		"opening_denominations": [dict(row, amount=row["denomination"] * row["count"]) for row in counted],
		"opening_remarks": remarks,
	}).insert(ignore_permissions=True)

	return {"session": session.name, "counter": session.counter}


def get_own_open_session_doc():
	name = get_open_session()
	if not name:
		frappe.throw("You have no open day to close.")
	return frappe.get_doc("Counter Session", name)


def apply_count(session, counted_cash, denominations, expenditure, payment_totals):
	"""Put a till count -- the notes and coins (or the total), the cash paid out,
	and the non-cash totals -- on `session`, without saving it."""
	session.set("denominations", parse_denominations(denominations))
	session.set("payment_mode_totals", parse_payment_totals(payment_totals))

	if not session.denominations and counted_cash in (None, ""):
		frappe.throw("Enter the counted cash, or count the notes and coins.")

	if flt(expenditure) < 0:
		frappe.throw("Expenditure cannot be negative.")

	session.counted_cash = flt(counted_cash)
	session.cash_expenditure = flt(expenditure)


def count_and_close(session, counted_cash, denominations, remarks, closed_by, expenditure=None, payment_totals=None):
	"""Record the till count (and the cash paid out of the till, `expenditure`,
	and the non-cash totals, `payment_totals`) on `session` and close it.

	A clean count submits the session. Otherwise it stays in Closing for a
	supervisor -- unless a supervisor is closing someone else's day, when their
	close is its own approval. Either way the counter is free from here on.
	"""
	apply_count(session, counted_cash, denominations, expenditure, payment_totals)
	session.close_remarks = remarks
	session.status = "Closing"
	session.closed_on = now_datetime()
	session.closed_by = closed_by
	session.save(ignore_permissions=True)

	if (
		abs(flt(session.difference)) >= TOLERANCE
		and closed_by != session.cashier
		and is_supervisor(closed_by)
	):
		session.approved_by = closed_by
		session.save(ignore_permissions=True)

	if abs(flt(session.difference)) < TOLERANCE or session.approved_by:
		session.flags.ignore_permissions = True
		session.submit()
		return {"session": session.name, "status": "Closed", "difference": flt(session.difference)}

	return {"session": session.name, "status": "Closing", "difference": session.difference}


@frappe.whitelist()
def close_day(counted_cash=None, denominations=None, remarks=None, expenditure=None, payment_totals=None):
	"""Close the user's open day with the counted cash.

	A clean count closes (submits) the session. A count that differs from the
	expected cash leaves it in Closing until a supervisor approves it (approve_close).
	`denominations` is a list of {denomination, count}; when given, it makes up the
	counted cash. `payment_totals` is a list of {payment_mode, amount} for the
	non-cash modes, recorded for the supervisor's day verification.

	A blind close: the cashier is never told the expected cash, so the difference
	is not sent back.
	"""
	session = get_own_open_session_doc()
	if session.status != "Open":
		frappe.throw(f"{session.name} is already counted and waiting for approval.")

	result = count_and_close(
		session, counted_cash, denominations, remarks, frappe.session.user, expenditure, payment_totals
	)
	result.pop("difference", None)
	return result


@frappe.whitelist()
def supervisor_close(session, counted_cash=None, denominations=None, remarks=None, expenditure=None, payment_totals=None):
	"""A supervisor counts the till and closes an open day from the Day Close grid.

	For a cashier who left without closing, which would otherwise hold the counter,
	or the supervisor's own day. The supervisor is recorded as closed_by; closing
	another cashier's day approves any difference in the same step and needs a reason.
	"""
	if not is_supervisor():
		frappe.throw("Only a supervisor can close another cashier's day.", frappe.PermissionError)

	doc = frappe.get_doc("Counter Session", session)
	if doc.docstatus != 0 or doc.status != "Open":
		frappe.throw(f"{session} is not an open day.")

	if doc.cashier != frappe.session.user and not (remarks or "").strip():
		frappe.throw("Give a reason for closing this day on the cashier's behalf.")

	return count_and_close(doc, counted_cash, denominations, remarks, frappe.session.user, expenditure, payment_totals)


@frappe.whitelist()
def approve_close(session):
	"""A supervisor signs off a close with a difference, which closes the session."""
	if not is_supervisor():
		frappe.throw("Only a supervisor can approve a cash difference.", frappe.PermissionError)

	doc = frappe.get_doc("Counter Session", session)
	if doc.docstatus != 0 or doc.status != "Closing":
		frappe.throw(f"{session} is not waiting for approval.")

	doc.approved_by = frappe.session.user
	doc.save(ignore_permissions=True)
	doc.flags.ignore_permissions = True
	doc.submit()
	return {"session": doc.name, "status": doc.status, "difference": doc.difference}


# ------------------------------------------------------------ supervisor's day close
# A supervisor's Day Close is not a blind count: it lays every counter's day for a
# date side by side -- what the system took and what the cashier entered -- with
# the totals last. The same lines feed the page, the Day Close Summary print and
# the Excel export (day_overview_lines), so the three always agree.


def build_day_overview(date=None):
	from digitz_erp.digitz_erp.doctype.counter_session.counter_session import get_takings_by_mode

	date = getdate(date or nowdate())
	is_today = date == getdate(nowdate())
	# The date's days, and -- looking at today -- any day still unfinished from before,
	# which still has to be counted
	names = frappe.db.sql_list(
		"""
		SELECT name FROM `tabCounter Session`
		WHERE docstatus < 2 AND (DATE(opened_on) = %(date)s OR (%(is_today)s AND docstatus = 0))
		ORDER BY counter, opened_on
		""",
		{"date": date, "is_today": 1 if is_today else 0},
	)
	docs = [frappe.get_doc("Counter Session", name) for name in names]
	modes = get_non_cash_payment_modes()
	system = get_takings_by_mode({doc.name: doc.source_session() for doc in docs})

	sessions = []
	for doc in docs:
		if doc.docstatus == 0:
			doc.calculate_cash()  # live figures for a day still open or waiting
		counted = doc.status != "Open"
		entered = {row.payment_mode: flt(row.amount) for row in doc.payment_mode_totals}
		sessions.append(frappe._dict({
			"name": doc.name,
			"counter": doc.counter,
			"cashier": doc.cashier,
			"cashier_name": frappe.utils.get_fullname(doc.cashier),
			"status": doc.status,
			"docstatus": doc.docstatus,
			"opened_on": doc.opened_on,
			"closed_on": doc.closed_on,
			"carried_over": getdate(doc.opened_on) != date,
			"amended_from": doc.amended_from,
			"closed_by_name": frappe.utils.get_fullname(doc.closed_by) if doc.closed_by else None,
			"approved_by_name": frappe.utils.get_fullname(doc.approved_by) if doc.approved_by else None,
			"close_remarks": doc.close_remarks,
			"is_own": doc.cashier == frappe.session.user,
			"denominations": [{"denomination": r.denomination, "count": r.count} for r in doc.denominations],
			"opening_float": flt(doc.opening_float),
			"cash_sales": flt(doc.cash_sales),
			"cash_receipts": flt(doc.cash_receipts),
			"cash_refunds": flt(doc.cash_refunds),
			"cash_paid_out": flt(doc.cash_paid_out),
			"cash_expenditure": flt(doc.cash_expenditure),
			"expected_cash": flt(doc.expected_cash),
			# Not counted yet: nothing to show, rather than a misleading 0
			"counted_cash": flt(doc.counted_cash) if counted else None,
			"difference": flt(doc.difference) if counted else None,
			"modes": {
				m.payment_mode: {
					"entered": entered.get(m.payment_mode) if counted else None,
					"system": flt(system[doc.name].get(m.payment_mode)),
				}
				for m in modes
			},
		}))

	overview = frappe._dict({
		"date": date,
		"is_today": is_today,
		"modes": modes,
		"sessions": sessions,
		"denominations": DENOMINATIONS,
	})
	overview.lines = day_overview_lines(overview)
	return overview


def day_overview_lines(overview):
	"""The rows of the grid: [{label, kind, values (one per session), total}].

	kind: "section" (a heading), "money", "strong", "difference" (coloured).
	A value of None is a figure not there yet (a day still open).
	"""
	def total(values):
		known = [v for v in values if v is not None]
		return sum(known) if known else None

	def line(label, values, kind="money"):
		return {"label": label, "kind": kind, "values": values, "total": total(values)}

	def field(name):
		return [s[name] for s in overview.sessions]

	def less(a, b):
		return None if a is None or b is None else a - b

	lines = [
		{"label": "Cash", "kind": "section"},
		line("Opening Float", field("opening_float")),
		line("+ Cash Sales", field("cash_sales")),
		line("+ Cash Receipts (Credit Collections)", field("cash_receipts")),
		line("− Cash Refunds", field("cash_refunds")),
		line("− Cash Paid Out", field("cash_paid_out")),
		line("− Expenditure (entered)", field("cash_expenditure")),
		line("Expected Cash", field("expected_cash"), "strong"),
		line("Counted Cash (entered)", field("counted_cash"), "strong"),
		line("Cash Difference", field("difference"), "difference"),
	]
	for m in overview.modes:
		mode = m.payment_mode
		entered = [s.modes[mode]["entered"] for s in overview.sessions]
		system = [s.modes[mode]["system"] for s in overview.sessions]
		lines += [
			{"label": f"{mode} ({m.mode})", "kind": "section"},
			line(f"{mode}: System", system),
			line(f"{mode}: Entered", entered, "strong"),
			line(f"{mode}: Difference", [less(e, sy) for e, sy in zip(entered, system)], "difference"),
		]
	return lines


def get_slip_payment_rows(doc):
	"""Day Close Slip: the non-cash lines of `doc`. A supervisor gets the system
	figure and the difference too; a cashier only what they entered."""
	from digitz_erp.digitz_erp.doctype.counter_session.counter_session import get_takings_by_mode

	entered = {row.payment_mode: flt(row.amount) for row in doc.payment_mode_totals}
	if not is_supervisor():
		return [{"payment_mode": mode, "entered": amount} for mode, amount in entered.items()]

	system = get_takings_by_mode({doc.name: doc.source_session()})[doc.name]
	rows = []
	for m in get_non_cash_payment_modes():
		mode = m.payment_mode
		if mode not in entered and not system.get(mode):
			continue
		row = {"payment_mode": mode, "system": flt(system.get(mode)), "entered": entered.get(mode)}
		row["difference"] = None if row["entered"] is None else row["entered"] - row["system"]
		rows.append(row)
	return rows


@frappe.whitelist()
def get_day_overview(date=None):
	"""Every counter's day for `date` (default today), for the supervisor's Day Close."""
	if not is_supervisor():
		frappe.throw("Only a supervisor can see every counter's day.", frappe.PermissionError)
	return build_day_overview(date)


@frappe.whitelist()
def get_day_summary_html(date=None):
	"""The Day Close Summary print: the grid of a date, landscape, with signatures."""
	if not is_supervisor():
		frappe.throw("Only a supervisor can print the Day Close Summary.", frappe.PermissionError)
	overview = build_day_overview(date)
	default_company = frappe.db.get_single_value("Global Settings", "default_company")
	return frappe.render_template(
		"digitz_erp/templates/day_close_summary.html",
		{
			"o": overview,
			"company": default_company,
			"currency": frappe.db.get_value("Company", default_company, "default_currency") if default_company else None,
			"printed_by": frappe.utils.get_fullname(frappe.session.user),
			"printed_on": now_datetime(),
		},
	)


@frappe.whitelist()
def export_day_overview(date=None):
	"""The Day Close grid of a date as a formatted Excel workbook: System Manager only.

	Laid out like the Print Summary -- one column per session, the totals last --
	with the header and first column frozen, and set to print landscape on one
	page wide.
	"""
	frappe.only_for("System Manager")

	from io import BytesIO

	from openpyxl import Workbook
	from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
	from openpyxl.utils import get_column_letter

	o = build_day_overview(date)
	default_company = frappe.db.get_single_value("Global Settings", "default_company")
	sessions = o.sessions
	last_col = len(sessions) + 2  # label column, the sessions, the total

	navy, blue, grey, light = "1F3A5F", "DCE6F1", "F2F2F2", "F7F9FC"
	thin = Side(style="thin", color="BFBFBF")
	box = Border(left=thin, right=thin, top=thin, bottom=thin)
	money_fmt = '#,##0.00;-#,##0.00;"-"'
	status_fill = {"Open": "D9F2E3", "Closing": "FCE8C8", "Closed": "DCE6F1"}

	def fill(colour):
		return PatternFill("solid", start_color=colour, end_color=colour)

	wb = Workbook()
	ws = wb.active
	ws.title = "Day Close"
	ws.sheet_view.showGridLines = False

	def merged(row, text, font, row_fill=None, height=None):
		ws.merge_cells(start_row=row, start_column=1, end_row=row, end_column=last_col)
		cell = ws.cell(row=row, column=1, value=text)
		cell.font = font
		cell.alignment = Alignment(horizontal="center", vertical="center")
		if row_fill:
			for col in range(1, last_col + 1):
				ws.cell(row=row, column=col).fill = fill(row_fill)
		if height:
			ws.row_dimensions[row].height = height

	merged(1, "DAY CLOSE SUMMARY", Font(size=16, bold=True, color="FFFFFF"), navy, 30)
	merged(2, " · ".join(filter(None, [
		default_company,
		f"Date: {frappe.utils.formatdate(o.date)}",
		f"Exported {frappe.utils.format_datetime(now_datetime())} by {frappe.utils.get_fullname(frappe.session.user)}",
	])), Font(size=10, italic=True, color="595959"), light, 20)

	if not sessions:
		merged(4, "No day was opened on any counter on this date.", Font(size=11, italic=True, color="7F7F7F"))
	else:
		def time(dt):
			return frappe.utils.get_datetime(dt).strftime("%I:%M %p") if dt else "open"

		# The column heads: counter and cashier, then the session, its hours and status
		heads = [
			("", [f"{s.counter}\n{s.cashier_name}" for s in sessions], "TOTAL"),
			("Session", [s.name + (f" (corrects {s.amended_from})" if s.amended_from else "") for s in sessions], ""),
			("Opened → Closed", [f"{time(s.opened_on)} → {time(s.closed_on)}" for s in sessions], ""),
			("Status", ["Awaiting Approval" if s.status == "Closing" else s.status for s in sessions], ""),
		]
		row = 4
		for i, (label, values, total) in enumerate(heads):
			for col, value in enumerate([label, *values, total], start=1):
				cell = ws.cell(row=row, column=col, value=value)
				cell.border = box
				cell.alignment = Alignment(horizontal="left" if col == 1 else "center", vertical="center", wrap_text=True)
				if i == 0:
					cell.font = Font(bold=True, color="FFFFFF", size=11)
					cell.fill = fill(navy)
				else:
					cell.font = Font(size=9, color="404040", bold=col == 1)
					cell.fill = fill(status_fill.get(sessions[col - 2].status, light) if i == 3 and 1 < col < last_col else light)
			row += 1
		ws.row_dimensions[4].height = 32
		if any(s.amended_from for s in sessions):
			ws.row_dimensions[5].height = 26  # room for "(corrects ...)"

		for ln in o.lines:
			if ln["kind"] == "section":
				merged(row, ln["label"].upper(), Font(bold=True, size=9, color=navy), blue)
				ws.cell(row=row, column=1).alignment = Alignment(horizontal="left", vertical="center")
				for col in range(1, last_col + 1):
					ws.cell(row=row, column=col).border = box
				row += 1
				continue

			bold = ln["kind"] in ("strong", "difference")
			for col, value in enumerate([ln["label"], *ln["values"], ln["total"]], start=1):
				cell = ws.cell(row=row, column=col)
				cell.border = box
				if col == 1:
					cell.value = ln["label"]
					cell.font = Font(bold=bold)
					continue
				if value is None:
					# Not counted yet: a dash, not a misleading zero
					cell.value = "—"
					cell.font = Font(color="A6A6A6")
					cell.alignment = Alignment(horizontal="center")
				else:
					cell.value = flt(value, 2)
					cell.number_format = money_fmt
					colour = None
					if ln["kind"] == "difference" and abs(flt(value)) >= TOLERANCE:
						colour = "C00000" if value < 0 else "C65911"
					cell.font = Font(bold=bold or col == last_col, color=colour)
				if col == last_col:
					cell.fill = fill(grey)
			if ln["kind"] == "strong":
				for col in range(1, last_col + 1):
					ws.cell(row=row, column=col).border = Border(left=thin, right=thin, top=Side(style="medium", color="404040"), bottom=thin)
			row += 1

		# Remarks, wrapped under each column
		ws.cell(row=row, column=1, value="Remarks").font = Font(bold=True)
		for col, s in enumerate(sessions, start=2):
			note = "\n".join(filter(None, [s.close_remarks, f"Approved by {s.approved_by_name}" if s.approved_by_name else None]))
			ws.cell(row=row, column=col, value=note or "")
		for col in range(1, last_col + 1):
			cell = ws.cell(row=row, column=col)
			cell.border = box
			cell.alignment = Alignment(vertical="top", wrap_text=True)
			if col > 1:
				cell.font = Font(size=9, color="404040")
		ws.row_dimensions[row].height = 45
		row += 2

		note = ws.cell(row=row, column=1, value="System: what the submitted documents add up to. Entered: what the cashier counted or typed at Day Close.")
		note.font = Font(size=8, italic=True, color="7F7F7F")
		row += 3

		# Signature lines, spread across the sheet
		spots = [1, max(2, last_col // 2 + 1), last_col] if last_col >= 3 else [1, 2]
		for col, label in zip(spots, ["Prepared by", "Checked by", "Approved by"]):
			cell = ws.cell(row=row, column=col, value=label)
			cell.border = Border(top=Side(style="thin", color="000000"))
			cell.alignment = Alignment(horizontal="center")
			cell.font = Font(size=9)

		ws.freeze_panes = "B5"
		ws.print_title_rows = "4:7"

	ws.column_dimensions["A"].width = 38
	for col in range(2, last_col + 1):
		ws.column_dimensions[get_column_letter(col)].width = 22
	ws.page_setup.orientation = "landscape"
	ws.page_setup.paperSize = ws.PAPERSIZE_A4
	ws.page_setup.fitToWidth = 1
	ws.page_setup.fitToHeight = 0
	ws.sheet_properties.pageSetUpPr.fitToPage = True
	ws.print_options.horizontalCentered = True

	out = BytesIO()
	wb.save(out)
	frappe.response["filename"] = f"Day Close {o.date}.xlsx"
	frappe.response["filecontent"] = out.getvalue()
	frappe.response["type"] = "binary"


@frappe.whitelist()
def edit_close(session, counted_cash=None, denominations=None, remarks=None, expenditure=None, payment_totals=None):
	"""A supervisor corrects a counted day: the count, the expenditure and the
	non-cash totals. Their correction is its own approval, so the day closes.

	A close still waiting for approval is corrected in place. A closed day is
	reopened: it is cancelled and an amended copy takes the corrected count, so
	the original stays on record. The cashier's day stays closed either way.
	"""
	if not is_supervisor():
		frappe.throw("Only a supervisor can correct a closed day.", frappe.PermissionError)
	if not (remarks or "").strip():
		frappe.throw("Give a reason for the correction.")

	doc = frappe.get_doc("Counter Session", session)
	if doc.docstatus == 0 and doc.status == "Closing":
		target = doc
	elif doc.docstatus == 1:
		# The documents of the day link to this session; they stay with it
		doc.flags.ignore_links = True
		doc.flags.ignore_permissions = True
		doc.cancel()
		target = frappe.copy_doc(doc)
		target.amended_from = doc.name
	else:
		frappe.throw(f"{session} is not counted yet: close it first.")

	apply_count(target, counted_cash, denominations, expenditure, payment_totals)
	note = f"Corrected by {frappe.utils.get_fullname(frappe.session.user)} on {frappe.utils.format_datetime(now_datetime())}: {remarks.strip()}"
	target.close_remarks = "\n".join(r for r in (doc.close_remarks, note) if r)
	target.status = "Closing"
	target.approved_by = frappe.session.user
	target.save(ignore_permissions=True)
	target.flags.ignore_permissions = True
	target.submit()
	return {"session": target.name, "status": target.status}


# The longest range Day History loads in one request
HISTORY_MAX_DAYS = 366


@frappe.whitelist()
def get_session_history(from_date=None, to_date=None, counter=None, cashier=None, status=None):
	"""Days for the Cashier Console's Day History: the user's own, or everyone's for a supervisor.
	A cashier gets only their own inputs (float, count, expenditure, non-cash totals).

	Reads through the Counter Session Summary report's query, which is raw SQL and
	checks no permissions, so the cashier filter is forced here: a cashier sees
	only their own days, as the doctype's Cashier permission (if_owner) allows.
	"""
	from digitz_erp.digitz_erp.report.counter_session_summary.counter_session_summary import get_data

	to_date = getdate(to_date or nowdate())
	from_date = getdate(from_date or add_days(to_date, -29))
	if from_date > to_date:
		from_date, to_date = to_date, from_date
	if date_diff(to_date, from_date) > HISTORY_MAX_DAYS:
		from_date = add_days(to_date, -HISTORY_MAX_DAYS)

	supervisor = is_supervisor()
	filters = {"from_date": from_date, "to_date": to_date}
	if status in ("Open", "Closing", "Closed"):
		filters["status"] = status
	if supervisor:
		if counter:
			filters["counter"] = counter
		if cashier:
			filters["cashier"] = cashier
	else:
		filters["cashier"] = frappe.session.user

	rows = get_data(filters)

	# What the cashier entered for card, bank and other modes, per day
	non_cash = dict(frappe.db.sql(
		"""SELECT parent, SUM(amount) FROM `tabCounter Session Payment Total`
		WHERE parenttype = 'Counter Session' AND parent IN %(names)s GROUP BY parent""",
		{"names": [r.name for r in rows] or [""]},
	))
	for r in rows:
		r.non_cash_total = flt(non_cash.get(r.name))

	totals = {
		"days": len(rows),
		"counted_cash": sum(flt(r.counted_cash) for r in rows if r.status != "Open"),
		"non_cash_total": sum(r.non_cash_total for r in rows),
		"cash_expenditure": sum(flt(r.cash_expenditure) for r in rows),
		"pending": sum(1 for r in rows if r.status in ("Open", "Closing")),
	}
	if supervisor:
		totals.update({
			"cash_sales": sum(flt(r.cash_sales) for r in rows),
			"cash_receipts": sum(flt(r.cash_receipts) for r in rows),
			"net_difference": sum(flt(r.difference) for r in rows if r.status != "Open"),
		})
	else:
		# A cashier sees only what they entered, never the day's sales or the
		# expected cash (a blind close), so those are not sent at all
		for r in rows:
			for field in SUPERVISOR_ONLY_FIELDS:
				r.pop(field, None)

	return {
		"rows": rows,
		"is_supervisor": supervisor,
		"from_date": from_date,
		"to_date": to_date,
		"totals": totals,
	}
